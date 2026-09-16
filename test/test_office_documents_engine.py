"""Tests for the offline OOXML document engine (W07 · office_documents).

Covers structured read, template creation, targeted in-place edit with
byte-preserving write-back, independent-parser reopen verification, and the
protected/unsupported-format rejection gate with a negative test per category.

Fidelity is asserted two ways: (1) untouched parts survive a round-trip
byte-for-byte, and (2) the written file reopens under a DIFFERENT reader than
the one that wrote it (python-docx for docx; a from-scratch stdlib zip walk for
pptx) and yields the expected content — so the test does not merely re-run the
engine's own code path over its own output.
"""

from __future__ import annotations

import os
import zipfile

# python-docx is a declared dependency and serves as the INDEPENDENT reader
# that reopens docx output the engine wrote through its own part-level path.
import docx as pydocx  # type: ignore[import-untyped]
import pytest

from kiro_crew.connections.vendors.microsoft import office_documents as od
from kiro_crew.connections.vendors.microsoft.office_documents import constants as C
from kiro_crew.connections.vendors.microsoft.office_documents import (
    container,
)
from kiro_crew.connections.vendors.microsoft.office_documents import docx as docx_mod
from kiro_crew.connections.vendors.microsoft.office_documents import pptx as pptx_mod
from kiro_crew.connections.vendors.microsoft.office_documents import (
    rejection,
)
from kiro_crew.connections.vendors.microsoft.office_documents.errors import (
    DocumentEditError,
    MalformedDocument,
    OfficeDocumentError,
    ProtectedDocument,
    UnsupportedDocument,
)

# ── Fixtures / builders ──


def _make_real_docx(path: str) -> None:
    """Author a real .docx (styles, theme, table) with python-docx."""
    d = pydocx.Document()
    d.add_heading("Original Title", level=1)
    d.add_paragraph("Body paragraph one.")
    d.add_paragraph("Body paragraph two.")
    table = d.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "r0c0"
    table.cell(0, 1).text = "r0c1"
    table.cell(1, 0).text = "r1c0"
    table.cell(1, 1).text = "r1c1"
    d.save(path)


_PPTX_SLIDE = (
    "<?xml version='1.0' encoding='UTF-8' standalone='yes'?>"
    "<p:sld xmlns:a='%s' xmlns:p='%s'>"
    "<p:cSld><p:spTree><p:sp><p:txBody>"
    "<a:p><a:r><a:t>{text}</a:t></a:r></a:p>"
    "</p:txBody></p:sp></p:spTree></p:cSld></p:sld>" % (C.A_NS, C.P_NS)
)
_PPTX_NOTES = (
    "<?xml version='1.0' encoding='UTF-8' standalone='yes'?>"
    "<p:notes xmlns:a='%s' xmlns:p='%s'>"
    "<p:cSld><p:spTree><p:sp><p:txBody>"
    "<a:p><a:r><a:t>{notes}</a:t></a:r></a:p>"
    "</p:txBody></p:sp></p:spTree></p:cSld></p:notes>" % (C.A_NS, C.P_NS)
)
_PPTX_SLIDE_RELS = (
    "<?xml version='1.0'?>"
    "<Relationships xmlns='http://schemas.openxmlformats.org/package/2006/relationships'>"
    "<Relationship Id='rId1' "
    "Type='http://schemas.openxmlformats.org/officeDocument/2006/relationships/notesSlide' "
    "Target='../notesSlides/notesSlide{n}.xml'/></Relationships>"
)


def _make_real_pptx(path: str, slides: list[tuple[str, str | None]]) -> None:
    """Author a .pptx with (body, notes) per slide. notes=None -> no notes part."""
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr(
            C.CONTENT_TYPES_PART,
            "<?xml version='1.0'?><Types xmlns='%s'/>" % C.CT_NS,
        )
        z.writestr(
            C.PPTX_PRESENTATION_PART,
            "<?xml version='1.0'?><p:presentation xmlns:p='%s'/>" % C.P_NS,
        )
        z.writestr("ppt/theme/theme1.xml", "<theme>original</theme>")
        for i, (body, notes) in enumerate(slides, 1):
            z.writestr(f"ppt/slides/slide{i}.xml", _PPTX_SLIDE.format(text=body))
            if notes is not None:
                z.writestr(
                    f"ppt/slides/_rels/slide{i}.xml.rels",
                    _PPTX_SLIDE_RELS.format(n=i),
                )
                z.writestr(
                    f"ppt/notesSlides/notesSlide{i}.xml",
                    _PPTX_NOTES.format(notes=notes),
                )


def _slide_text_via_stdlib(path: str, n: int) -> str:
    """Independent pptx slide-text reader: bare stdlib, not the engine's code."""
    # nosemgrep: python.lang.security.use-defused-xml.use-defused-xml
    # This helper is the INDEPENDENT verification reader: it deliberately
    # parses engine output via stdlib (not defusedxml) to prove the output is
    # interoperable with a different parser. The input is engine-produced (safe).
    import xml.etree.ElementTree as ET

    with zipfile.ZipFile(path) as z:
        raw = z.read(f"ppt/slides/slide{n}.xml")
    root = ET.fromstring(raw)
    return "".join(t.text or "" for t in root.iter(f"{C.A}t"))


@pytest.fixture()
def docx_path(tmp_path):
    p = str(tmp_path / "doc.docx")
    _make_real_docx(p)
    return p


@pytest.fixture()
def pptx_path(tmp_path):
    p = str(tmp_path / "deck.pptx")
    _make_real_pptx(p, [("Slide one body", "Note one"), ("Slide two body", None)])
    return p


# ── DOCX structured read ──


class TestDocxRead:
    def test_paragraphs_and_styles(self, docx_path):
        content = od.read_document(docx_path)
        texts = [p.text for p in content.paragraphs]
        assert "Original Title" in texts
        assert "Body paragraph one." in texts
        assert content.paragraphs[0].style.startswith("Heading")

    def test_tables(self, docx_path):
        content = od.read_document(docx_path)
        assert len(content.tables) == 1
        assert content.tables[0].rows == [["r0c0", "r0c1"], ["r1c0", "r1c1"]]

    def test_text_property_excludes_tables(self, docx_path):
        content = od.read_document(docx_path)
        assert "Body paragraph one." in content.text
        assert "r0c0" not in content.text

    def test_facade_read_dispatches_docx(self, docx_path):
        content = od.read(docx_path)
        assert isinstance(content, od.DocxContent)

    def test_empty_body_returns_empty(self, tmp_path):
        p = str(tmp_path / "empty.docx")
        with zipfile.ZipFile(p, "w") as z:
            z.writestr(
                C.DOCX_MAIN_PART,
                "<?xml version='1.0'?><w:document xmlns:w='%s'><w:body/></w:document>" % C.W_NS,
            )
        content = od.read_document(p)
        assert content.paragraphs == []
        assert content.tables == []

    def test_document_without_body(self, tmp_path):
        p = str(tmp_path / "nobody.docx")
        with zipfile.ZipFile(p, "w") as z:
            z.writestr(
                C.DOCX_MAIN_PART,
                "<?xml version='1.0'?><w:document xmlns:w='%s'/>" % C.W_NS,
            )
        assert od.read_document(p).paragraphs == []


# ── DOCX in-place edit + independent reopen + fidelity ──


class TestDocxEdit:
    def test_edit_reopens_under_independent_parser(self, docx_path, tmp_path):
        out = str(tmp_path / "out.docx")
        od.edit_in_place(docx_path, out, {1: "Rewritten body."})
        reopened = [p.text for p in pydocx.Document(out).paragraphs]
        assert "Rewritten body." in reopened
        assert "Body paragraph two." in reopened  # untouched paragraph survives
        assert "Original Title" in reopened

    def test_untouched_parts_preserved_byte_for_byte(self, docx_path, tmp_path):
        out = str(tmp_path / "out.docx")
        src_zip = zipfile.ZipFile(docx_path)
        before = {n: src_zip.read(n) for n in src_zip.namelist() if n != C.DOCX_MAIN_PART}
        od.edit_in_place(docx_path, out, {1: "Changed."})
        out_zip = zipfile.ZipFile(out)
        for name, data in before.items():
            assert out_zip.read(name) == data, f"part {name} not preserved"

    def test_style_preserved_on_edited_paragraph(self, docx_path, tmp_path):
        out = str(tmp_path / "out.docx")
        od.edit_in_place(docx_path, out, {0: "New Heading Text"})
        reopened = pydocx.Document(out)
        assert reopened.paragraphs[0].text == "New Heading Text"
        assert reopened.paragraphs[0].style.name.startswith("Heading")

    def test_in_place_same_path(self, docx_path):
        od.edit_in_place(docx_path, docx_path, {1: "In place."})
        assert "In place." in [p.text for p in pydocx.Document(docx_path).paragraphs]

    def test_edit_out_of_range_raises_before_write(self, docx_path, tmp_path):
        out = str(tmp_path / "out.docx")
        with pytest.raises(DocumentEditError):
            od.edit_in_place(docx_path, out, {99: "nope"})
        assert not os.path.exists(out)

    def test_edit_illegal_xml_char_raises_before_write(self, docx_path, tmp_path):
        # A NUL (XML-1.0-forbidden) in replacement text must be refused before
        # any bytes are written, not serialised into unreopenable output.
        out = str(tmp_path / "out.docx")
        with pytest.raises(DocumentEditError):
            od.edit_in_place(docx_path, out, {0: "bad\x00text"})
        assert not os.path.exists(out)

    def test_no_op_edit_produces_faithful_copy(self, docx_path, tmp_path):
        out = str(tmp_path / "copy.docx")
        docx_mod.replace_paragraph_text(docx_path, out, {})
        assert [p.text for p in pydocx.Document(out).paragraphs] == [
            p.text for p in pydocx.Document(docx_path).paragraphs
        ]

    def test_whitespace_preserved(self, docx_path, tmp_path):
        out = str(tmp_path / "ws.docx")
        od.edit_in_place(docx_path, out, {1: "  leading and trailing  "})
        assert "  leading and trailing  " in [p.text for p in pydocx.Document(out).paragraphs]


# ── DOCX template create ──


class TestDocxTemplate:
    def test_create_carries_parts_and_fills(self, docx_path, tmp_path):
        out = str(tmp_path / "fromtpl.docx")
        kind = od.create_from_template(docx_path, out, {0: "Filled Title"})
        assert kind == "docx"
        reopened = pydocx.Document(out)
        assert reopened.paragraphs[0].text == "Filled Title"
        # Template's styles part carried over.
        assert "word/styles.xml" in zipfile.ZipFile(out).namelist()

    def test_create_without_edits_is_faithful_copy(self, docx_path, tmp_path):
        out = str(tmp_path / "copy.docx")
        od.create_from_template(docx_path, out)
        assert [p.text for p in pydocx.Document(out).paragraphs] == [
            p.text for p in pydocx.Document(docx_path).paragraphs
        ]


# ── PPTX structured read (stdlib) ──


class TestPptxRead:
    def test_slides_and_notes(self, pptx_path):
        content = od.read_presentation(pptx_path)
        assert [s.number for s in content.slides] == [1, 2]
        assert content.slides[0].text == "Slide one body"
        assert content.slides[0].notes == "Note one"
        assert content.slides[1].notes == ""  # no notes part

    def test_facade_read_dispatches_pptx(self, pptx_path):
        assert isinstance(od.read(pptx_path), od.PptxContent)

    def test_text_property(self, pptx_path):
        content = od.read_presentation(pptx_path)
        assert "Slide one body" in content.text
        assert "Slide two body" in content.text

    def test_slides_sorted_numerically(self, tmp_path):
        # slide10 must sort after slide2, not lexically before it.
        p = str(tmp_path / "many.pptx")
        _make_real_pptx(p, [(f"body{i}", None) for i in range(1, 11)])
        content = od.read_presentation(p)
        assert [s.number for s in content.slides] == list(range(1, 11))


# ── PPTX edit + independent reopen + fidelity ──


class TestPptxEdit:
    def test_edit_reopens_via_stdlib(self, pptx_path, tmp_path):
        out = str(tmp_path / "out.pptx")
        od.edit_in_place(pptx_path, out, {1: "Edited one"})
        assert _slide_text_via_stdlib(out, 1) == "Edited one"
        # Second slide untouched.
        assert _slide_text_via_stdlib(out, 2) == "Slide two body"

    def test_notes_and_other_parts_preserved(self, pptx_path, tmp_path):
        out = str(tmp_path / "out.pptx")
        src = zipfile.ZipFile(pptx_path)
        before = {n: src.read(n) for n in src.namelist() if n != "ppt/slides/slide1.xml"}
        od.edit_in_place(pptx_path, out, {1: "Edited"})
        out_zip = zipfile.ZipFile(out)
        for name, data in before.items():
            assert out_zip.read(name) == data, f"part {name} not preserved"
        # Notes still readable and unchanged via the engine too.
        assert od.read_presentation(out).slides[0].notes == "Note one"
        # Theme preserved.
        assert out_zip.read("ppt/theme/theme1.xml") == b"<theme>original</theme>"

    def test_unknown_slide_raises_before_write(self, pptx_path, tmp_path):
        out = str(tmp_path / "out.pptx")
        with pytest.raises(DocumentEditError):
            od.edit_in_place(pptx_path, out, {9: "nope"})
        assert not os.path.exists(out)

    def test_edit_illegal_xml_char_raises_before_write(self, pptx_path, tmp_path):
        out = str(tmp_path / "out.pptx")
        with pytest.raises(DocumentEditError):
            od.edit_in_place(pptx_path, out, {1: "bad\x0bctrl"})
        assert not os.path.exists(out)

    def test_slide_with_no_run_raises(self, tmp_path):
        p = str(tmp_path / "empty_slide.pptx")
        with zipfile.ZipFile(p, "w") as z:
            z.writestr(C.PPTX_PRESENTATION_PART, "<p:presentation xmlns:p='%s'/>" % C.P_NS)
            z.writestr(
                "ppt/slides/slide1.xml",
                "<p:sld xmlns:a='%s' xmlns:p='%s'><p:cSld><p:spTree/></p:cSld></p:sld>"
                % (C.A_NS, C.P_NS),
            )
        with pytest.raises(DocumentEditError):
            od.edit_in_place(p, str(tmp_path / "o.pptx"), {1: "x"})

    def test_no_op_edit_faithful_copy(self, pptx_path, tmp_path):
        out = str(tmp_path / "copy.pptx")
        pptx_mod.replace_slide_text(pptx_path, out, {})
        assert _slide_text_via_stdlib(out, 1) == "Slide one body"

    def test_create_from_template(self, pptx_path, tmp_path):
        out = str(tmp_path / "tpl.pptx")
        kind = od.create_from_template(pptx_path, out, {1: "Templated"})
        assert kind == "pptx"
        assert _slide_text_via_stdlib(out, 1) == "Templated"


# ── Rejection gate: one negative test per category ──


class TestRejectionGate:
    def _write(self, tmp_path, name, data: bytes) -> str:
        p = str(tmp_path / name)
        with open(p, "wb") as fh:
            fh.write(data)
        return p

    def test_legacy_ole2_doc(self, tmp_path):
        p = self._write(tmp_path, "old.doc", C.OLE2_MAGIC + b"\x00" * 64)
        v = od.classify(p)
        assert not v.ok and v.reason == UnsupportedDocument.reason
        with pytest.raises(UnsupportedDocument):
            rejection.ensure_editable(p)

    def test_encrypted_ooxml_wrapper(self, tmp_path):
        p = self._write(tmp_path, "enc.docx", C.OLE2_MAGIC + b"\x00" * 64)
        v = od.classify(p)
        assert not v.ok and v.reason == ProtectedDocument.reason
        with pytest.raises(ProtectedDocument):
            rejection.ensure_editable(p)

    def test_macro_enabled_extension(self, tmp_path):
        p = str(tmp_path / "m.docm")
        with zipfile.ZipFile(p, "w") as z:
            z.writestr(C.DOCX_MAIN_PART, "<a/>")
        assert od.classify(p).reason == ProtectedDocument.reason

    def test_vba_project_regardless_of_extension(self, tmp_path):
        p = str(tmp_path / "v.docx")
        with zipfile.ZipFile(p, "w") as z:
            z.writestr(C.DOCX_MAIN_PART, "<a/>")
            z.writestr(C.VBA_PROJECT_PART, b"macro")
        assert od.classify(p).reason == ProtectedDocument.reason

    def test_digital_signature_presence(self, tmp_path):
        p = str(tmp_path / "s.docx")
        with zipfile.ZipFile(p, "w") as z:
            z.writestr(C.DOCX_MAIN_PART, "<a/>")
            z.writestr(C.SIGNATURE_ORIGIN_PART, b"sig")
        v = od.classify(p)
        assert v.reason == ProtectedDocument.reason
        assert "validity" in v.detail  # presence-only, validity not asserted

    def test_irm_protection(self, tmp_path):
        p = str(tmp_path / "i.docx")
        with zipfile.ZipFile(p, "w") as z:
            z.writestr(C.DOCX_MAIN_PART, "<a/>")
            z.writestr("\x06DataSpaces/DataSpaceMap.xml", b"drm")
        assert od.classify(p).reason == ProtectedDocument.reason

    def test_encrypted_package_stream_member(self, tmp_path):
        p = str(tmp_path / "ep.docx")
        with zipfile.ZipFile(p, "w") as z:
            z.writestr(C.DOCX_MAIN_PART, "<a/>")
            z.writestr("EncryptedPackage", b"blob")
        assert od.classify(p).reason == ProtectedDocument.reason

    def test_not_a_zip(self, tmp_path):
        p = self._write(tmp_path, "g.docx", b"this is not a zip")
        v = od.classify(p)
        assert v.reason == MalformedDocument.reason
        with pytest.raises(MalformedDocument):
            rejection.ensure_editable(p)

    def test_empty_archive(self, tmp_path):
        p = self._write(tmp_path, "e.docx", C.ZIP_EMPTY_MAGIC + b"\x00" * 18)
        assert od.classify(p).reason == MalformedDocument.reason

    def test_zip_without_main_part(self, tmp_path):
        p = str(tmp_path / "x.docx")
        with zipfile.ZipFile(p, "w") as z:
            z.writestr("xl/workbook.xml", "<a/>")
        assert od.classify(p).reason == MalformedDocument.reason

    def test_missing_file(self, tmp_path):
        with pytest.raises(MalformedDocument):
            rejection.ensure_editable(str(tmp_path / "nope.docx"))

    def test_wrong_expected_kind(self, pptx_path):
        with pytest.raises(UnsupportedDocument):
            rejection.ensure_editable(pptx_path, expected_kind="docx")

    def test_read_document_rejects_protected(self, tmp_path):
        p = str(tmp_path / "v.docx")
        with zipfile.ZipFile(p, "w") as z:
            z.writestr(C.DOCX_MAIN_PART, "<a/>")
            z.writestr(C.VBA_PROJECT_PART, b"macro")
        with pytest.raises(ProtectedDocument):
            od.read_document(p)

    def test_legacy_extension_without_magic(self, tmp_path):
        # A .ppt name whose bytes are not OLE2 is still refused on extension.
        p = self._write(tmp_path, "x.ppt", b"garbage bytes")
        assert od.classify(p).reason == UnsupportedDocument.reason


# ── Container primitives ──


class TestContainer:
    def test_read_part_missing(self, docx_path):
        with pytest.raises(MalformedDocument):
            container.read_part(docx_path, "no/such/part.xml")

    def test_read_part_size_cap(self, tmp_path):
        p = str(tmp_path / "big.docx")
        with zipfile.ZipFile(p, "w") as z:
            z.writestr(C.DOCX_MAIN_PART, b"x" * 100)
        with pytest.raises(MalformedDocument):
            container.read_part(p, C.DOCX_MAIN_PART, max_size=10)

    def test_rewrite_replace_absent_part_raises(self, docx_path, tmp_path):
        out = str(tmp_path / "o.docx")
        with pytest.raises(MalformedDocument):
            container.rewrite_parts(docx_path, out, {"no/such.xml": b"x"})
        assert not os.path.exists(out)

    def test_rewrite_rejects_duplicate_member_names(self, docx_path, tmp_path):
        # A container that names one part twice is refused rather than copied
        # (the second entry would silently shadow the first for any reader).
        dup = str(tmp_path / "dup.docx")
        with zipfile.ZipFile(docx_path) as src, zipfile.ZipFile(dup, "w") as dst:
            for info in src.infolist():
                dst.writestr(info, src.read(info.filename))
            # Append a second entry with a name already present.
            dst.writestr("word/styles.xml", b"<duplicate/>")
        out = str(tmp_path / "o.docx")
        with pytest.raises(MalformedDocument):
            container.rewrite_parts(dup, out, {C.DOCX_MAIN_PART: b"<x/>"})
        assert not os.path.exists(out)

    def test_parse_xml_part(self, docx_path):
        root = container.parse_xml_part(docx_path, C.DOCX_MAIN_PART)
        assert root.tag == f"{C.W}document"

    def test_malformed_xml_part(self, tmp_path):
        p = str(tmp_path / "bad.docx")
        with zipfile.ZipFile(p, "w") as z:
            z.writestr(C.DOCX_MAIN_PART, b"<not well formed")
        with pytest.raises(MalformedDocument):
            container.parse_xml_part(p, C.DOCX_MAIN_PART)


# ── Error hierarchy ──


class TestConstants:
    def test_find_illegal_xml_char_passes_legal_text(self):
        assert C.find_illegal_xml_char("hello\tworld\r\nyay 🎉") is None

    def test_find_illegal_xml_char_flags_forbidden_control(self):
        assert C.find_illegal_xml_char("bad\x00nul") == "\x00"
        assert C.find_illegal_xml_char("vtab\x0bhere") == "\x0b"


class TestErrors:
    def test_protected_is_unsupported(self):
        assert issubclass(ProtectedDocument, UnsupportedDocument)
        assert issubclass(UnsupportedDocument, OfficeDocumentError)

    def test_reason_and_message(self):
        e = ProtectedDocument("locked file")
        assert e.reason == "protected_document"
        assert e.message == "locked file"

    def test_custom_reason(self):
        e = OfficeDocumentError("boom", reason="custom")
        assert e.reason == "custom"


# ── Parser-independence of the round-trip (fidelity matrix) ──


class TestFidelityMatrix:
    def test_docx_full_round_trip_all_parts(self, docx_path, tmp_path):
        """Every part except the edited one is byte-identical after edit."""
        out = str(tmp_path / "rt.docx")
        src = zipfile.ZipFile(docx_path)
        original = {n: src.read(n) for n in src.namelist()}
        od.edit_in_place(docx_path, out, {2: "Third paragraph changed."})
        result = zipfile.ZipFile(out)
        assert set(result.namelist()) == set(original)
        for name in original:
            if name == C.DOCX_MAIN_PART:
                continue
            assert result.read(name) == original[name]

    def test_pptx_round_trip_preserves_member_set(self, pptx_path, tmp_path):
        out = str(tmp_path / "rt.pptx")
        before = set(zipfile.ZipFile(pptx_path).namelist())
        od.edit_in_place(pptx_path, out, {1: "changed"})
        assert set(zipfile.ZipFile(out).namelist()) == before

    def test_double_round_trip_stable(self, docx_path, tmp_path):
        """Editing twice does not accrete drift in untouched parts."""
        out1 = str(tmp_path / "r1.docx")
        out2 = str(tmp_path / "r2.docx")
        od.edit_in_place(docx_path, out1, {1: "first"})
        od.edit_in_place(out1, out2, {1: "second"})
        z1, z2 = zipfile.ZipFile(out1), zipfile.ZipFile(out2)
        for name in z1.namelist():
            if name == C.DOCX_MAIN_PART:
                continue
            assert z1.read(name) == z2.read(name)
        assert "second" in [p.text for p in pydocx.Document(out2).paragraphs]


# ── Container edge paths ──


class TestContainerEdges:
    def test_sensitive_path_refused(self, docx_path, monkeypatch):
        monkeypatch.setattr(container, "is_sensitive_path", lambda p: True)
        with pytest.raises(OfficeDocumentError) as exc:
            container.read_part(docx_path, C.DOCX_MAIN_PART)
        assert exc.value.reason == "sensitive_path"

    def test_bad_zip_open_is_malformed(self, tmp_path):
        # A file that passes the vet's tail scan poorly / is not a real zip.
        p = str(tmp_path / "bad.docx")
        with open(p, "wb") as fh:
            fh.write(b"PK\x03\x04" + b"\x00" * 4 + b"corrupt")
        with pytest.raises(MalformedDocument):
            container.part_names(p)

    def test_rewrite_rejects_zip_bomb_member_on_copy(self, docx_path, tmp_path):
        # A container with one highly-compressible oversized member passes the
        # inventory vet (member count / CD bytes are small) but must be refused
        # on the byte-preserving copy path, not inflated into memory.
        bomb = str(tmp_path / "bomb.docx")
        big = b"\x00" * (C.MAX_PART_BYTES + 1024)
        with (
            zipfile.ZipFile(docx_path) as src,
            zipfile.ZipFile(bomb, "w", zipfile.ZIP_DEFLATED) as dst,
        ):
            for info in src.infolist():
                dst.writestr(info, src.read(info.filename))
            dst.writestr("word/media/bomb.bin", big)
        out = str(tmp_path / "o.docx")
        with pytest.raises(MalformedDocument):
            container.rewrite_parts(bomb, out, {C.DOCX_MAIN_PART: b"<x/>"})
        assert not os.path.exists(out)

    def test_rewrite_cleanup_on_write_failure(self, docx_path, tmp_path, monkeypatch):
        out = str(tmp_path / "o.docx")
        # Force the write to blow up mid-build; the temp artifact must be gone
        # and the destination untouched.
        real_writestr = zipfile.ZipFile.writestr

        def boom(self, *a, **k):
            raise RuntimeError("disk full")

        monkeypatch.setattr(zipfile.ZipFile, "writestr", boom)
        with pytest.raises(RuntimeError):
            container.rewrite_parts(docx_path, out, {C.DOCX_MAIN_PART: b"<x/>"})
        monkeypatch.setattr(zipfile.ZipFile, "writestr", real_writestr)
        assert not os.path.exists(out)
        # No leftover temp files in the destination dir.
        leftovers = [n for n in os.listdir(tmp_path) if n.startswith(".ooxml-")]
        assert leftovers == []

    def test_parse_without_defusedxml_refuses(self, docx_path, monkeypatch):
        monkeypatch.setattr(container, "_xml_fromstring", None)
        with pytest.raises(MalformedDocument):
            container.parse_xml_part(docx_path, C.DOCX_MAIN_PART)


# ── PPTX notes-resolution edge paths ──


class TestPptxNotesEdges:
    def test_notes_target_with_parent_ref(self, tmp_path):
        # Target uses ../notesSlides/ which must normalize correctly.
        p = str(tmp_path / "n.pptx")
        _make_real_pptx(p, [("body", "the note")])
        assert od.read_presentation(p).slides[0].notes == "the note"

    def test_notes_rel_pointing_at_missing_part(self, tmp_path):
        p = str(tmp_path / "dangling.pptx")
        with zipfile.ZipFile(p, "w") as z:
            z.writestr(C.PPTX_PRESENTATION_PART, "<p:presentation xmlns:p='%s'/>" % C.P_NS)
            z.writestr("ppt/slides/slide1.xml", _PPTX_SLIDE.format(text="hi"))
            z.writestr(
                "ppt/slides/_rels/slide1.xml.rels",
                _PPTX_SLIDE_RELS.format(n=1),  # points at notesSlide1 which is absent
            )
        # No notes part exists -> notes is empty, no crash.
        assert od.read_presentation(p).slides[0].notes == ""

    def test_replace_slide_no_slides_raises(self, tmp_path):
        p = str(tmp_path / "noslides.pptx")
        with zipfile.ZipFile(p, "w") as z:
            z.writestr(C.PPTX_PRESENTATION_PART, "<p:presentation xmlns:p='%s'/>" % C.P_NS)
        with pytest.raises(DocumentEditError):
            pptx_mod.replace_slide_text(p, str(tmp_path / "o.pptx"), {})

    def test_normalize_rel_target(self):
        assert (
            pptx_mod._normalize_rel_target("ppt/slides/", "../notesSlides/notesSlide1.xml")
            == "ppt/notesSlides/notesSlide1.xml"
        )
        assert (
            pptx_mod._normalize_rel_target("ppt/slides/", "./slideLayout1.xml")
            == "ppt/slides/slideLayout1.xml"
        )
        # A target that climbs above the package root names no real part.
        assert pptx_mod._normalize_rel_target("ppt/slides/", "../../../../etc") == ""
        assert pptx_mod._normalize_rel_target("ppt/slides/", "../..") == ""


# ── PPTX presentation order (sldIdLst, not filename number) ──


def _make_reordered_pptx(path: str, slide_bodies_in_file_order, sldid_order):
    """Author a .pptx where sldIdLst order differs from slideN.xml filename order.

    *slide_bodies_in_file_order* maps 1-based file number -> body text.
    *sldid_order* is the list of file numbers in the order the id list presents
    them (i.e. true presentation order).
    """
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr(C.CONTENT_TYPES_PART, "<?xml version='1.0'?><Types xmlns='%s'/>" % C.CT_NS)
        # presentation.xml carries a sldIdLst whose entries reference rIds.
        sld_ids = "".join(
            "<p:sldId id='%d' r:id='rId%d'/>" % (256 + i, n) for i, n in enumerate(sldid_order)
        )
        z.writestr(
            C.PPTX_PRESENTATION_PART,
            "<?xml version='1.0'?><p:presentation xmlns:p='%s' xmlns:r='%s'>"
            "<p:sldIdLst>%s</p:sldIdLst></p:presentation>" % (C.P_NS, C.R_NS, sld_ids),
        )
        # presentation.xml.rels maps each rId to its slide part.
        rels = "".join(
            "<Relationship Id='rId%d' "
            "Type='http://schemas.openxmlformats.org/officeDocument/2006/relationships/slide' "
            "Target='slides/slide%d.xml'/>" % (n, n)
            for n in sldid_order
        )
        z.writestr(
            "ppt/_rels/presentation.xml.rels",
            "<?xml version='1.0'?><Relationships "
            "xmlns='http://schemas.openxmlformats.org/package/2006/relationships'>%s"
            "</Relationships>" % rels,
        )
        for n, body in slide_bodies_in_file_order.items():
            z.writestr("ppt/slides/slide%d.xml" % n, _PPTX_SLIDE.format(text=body))


class TestPptxPresentationOrder:
    def test_read_follows_sldidlst_not_filename(self, tmp_path):
        # File slide1="A", slide2="B", but the id list presents 2 then 1.
        p = str(tmp_path / "reordered.pptx")
        _make_reordered_pptx(p, {1: "A", 2: "B"}, sldid_order=[2, 1])
        texts = [s.text for s in od.read_presentation(p).slides]
        assert texts == ["B", "A"]

    def test_edit_position_targets_presentation_order(self, tmp_path):
        # Editing position 1 must hit the slide the id list shows first (file 2).
        p = str(tmp_path / "reordered.pptx")
        _make_reordered_pptx(p, {1: "A", 2: "B"}, sldid_order=[2, 1])
        out = str(tmp_path / "o.pptx")
        od.edit_in_place(p, out, {1: "EDITED"})
        assert [s.text for s in od.read_presentation(out).slides] == ["EDITED", "A"]

    def test_dangling_sldid_falls_back_to_filename_order(self, tmp_path):
        # A sldIdLst rId with no matching rel -> fall back to filename order.
        p = str(tmp_path / "dangling.pptx")
        with zipfile.ZipFile(p, "w", zipfile.ZIP_DEFLATED) as z:
            z.writestr(C.CONTENT_TYPES_PART, "<?xml version='1.0'?><Types xmlns='%s'/>" % C.CT_NS)
            z.writestr(
                C.PPTX_PRESENTATION_PART,
                "<?xml version='1.0'?><p:presentation xmlns:p='%s' xmlns:r='%s'>"
                "<p:sldIdLst><p:sldId id='256' r:id='rNope'/></p:sldIdLst>"
                "</p:presentation>" % (C.P_NS, C.R_NS),
            )
            z.writestr(
                "ppt/_rels/presentation.xml.rels",
                "<?xml version='1.0'?><Relationships "
                "xmlns='http://schemas.openxmlformats.org/package/2006/relationships'/>",
            )
            z.writestr("ppt/slides/slide1.xml", _PPTX_SLIDE.format(text="only"))
        assert [s.text for s in od.read_presentation(p).slides] == ["only"]


class TestSensitivePathRejection:
    def test_classify_refuses_sensitive_path(self, tmp_path, monkeypatch):
        p = str(tmp_path / "secret.docx")
        with open(p, "wb") as fh:
            fh.write(b"PK\x03\x04ignored")
        monkeypatch.setattr(
            "kiro_crew.connections.vendors.microsoft.office_documents.rejection."
            "is_sensitive_path",
            lambda path: True,
        )
        verdict = rejection.classify(p)
        assert verdict.ok is False
        assert verdict.reason == "sensitive_path"

    def test_ensure_editable_raises_on_sensitive_path(self, tmp_path, monkeypatch):
        p = str(tmp_path / "secret.docx")
        with open(p, "wb") as fh:
            fh.write(b"PK\x03\x04ignored")
        monkeypatch.setattr(
            "kiro_crew.connections.vendors.microsoft.office_documents.rejection."
            "is_sensitive_path",
            lambda path: True,
        )
        with pytest.raises(OfficeDocumentError):
            rejection.ensure_editable(p)
